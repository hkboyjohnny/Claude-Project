"""Build Sinochem_Peer_Comparison.xlsx - inputs on 'Data', every ratio is a live Excel formula."""
import json
import pathlib

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

ROOT = pathlib.Path(__file__).resolve().parent.parent
D = json.loads((ROOT / "data" / "dataset.json").read_text())
FIELDS = D["fields"]
ENT = {e["key"]: e for e in D["entities"]}

# ----------------------------------------------------------------------------------------------------------------------- styles
NAVY, TEAL, AMBER, GREY, LIGHT = "1F3A5F", "2E7D8C", "FFF2CC", "EDEDED", "F7F9FB"
f_title = Font(name="Calibri", size=15, bold=True, color=NAVY)
f_sub = Font(name="Calibri", size=10, italic=True, color="555555")
f_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
f_in = Font(name="Calibri", size=10, color="0000FF")          # hard-coded input
f_calc = Font(name="Calibri", size=10, color="000000")
f_bold = Font(name="Calibri", size=10, bold=True)
fill_hdr = PatternFill("solid", fgColor=NAVY)
fill_grp = PatternFill("solid", fgColor=TEAL)
fill_sin = PatternFill("solid", fgColor=AMBER)
fill_stat = PatternFill("solid", fgColor=GREY)
fill_light = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()

# ----------------------------------------------------------------------------------------------------------------------- FX sheet
fx = wb.active
fx.title = "FX"
fx["A1"], fx["A1"].font = "Average exchange rates used for USD translation", f_title
fx["A2"], fx["A2"].font = "Local currency per 1 USD, calendar-2025 average (IRS yearly average rates). Edit the blue cells to re-translate every USD figure in the workbook.", f_sub
for j, h in enumerate(["Currency", "LC per USD (2025 avg)", "LC per USD (2024 avg)"], 1):
    c = fx.cell(row=4, column=j, value=h)
    c.font, c.fill, c.alignment, c.border = f_hdr, fill_hdr, ctr, box
for i, ccy in enumerate(["USD", "CNY", "EUR", "JPY"]):
    fx.cell(row=5 + i, column=1, value=ccy).border = box
    for j, rate in enumerate([D["fx_2025"][ccy], D["fx_2024"][ccy]], 2):
        c = fx.cell(row=5 + i, column=j, value=rate)
        c.font, c.number_format, c.border = f_in, "#,##0.000", box
fx.column_dimensions["A"].width, fx.column_dimensions["B"].width, fx.column_dimensions["C"].width = 14, 24, 24
fx["A11"] = "Source: https://www.irs.gov/individuals/international-taxpayers/yearly-average-currency-exchange-rates (tables for 2025 and 2024)."
fx["A11"].font = f_sub

# ----------------------------------------------------------------------------------------------------------------------- Data sheet
ds = wb.create_sheet("Data")
ds["A1"], ds["A1"].font = "Standardised financial inputs (millions of reporting currency)", f_title
ds["A2"] = ("Blue = hard-coded from filings; black = formula. EBIT = operating profit before financing as reported (China GAAP: operating profit + net finance expense); "
            "EBITDA = EBIT + D&A with no add-back for impairments/one-offs; Debt = borrowings + bonds + lease liabilities; Capex = cash paid for PP&E and intangibles.")
ds["A2"].font = f_sub
LABEL = {"revenue": "Revenue", "cogs": "Cost of sales", "gross_profit": "Gross profit", "ebit": "EBIT", "da": "D&A", "ebitda": "EBITDA",
         "interest_exp": "Interest expense", "pretax": "Pre-tax profit", "net_income": "Net income", "net_income_parent": "Net income (parent)",
         "total_assets": "Total assets", "total_liabilities": "Total liabilities", "total_equity": "Total equity", "parent_equity": "Parent equity",
         "cash": "Cash", "current_assets": "Current assets", "current_liabilities": "Current liabilities", "debt": "Debt (incl. leases)",
         "ocf": "Operating cash flow", "capex": "Capex", "tax_surcharge": "Taxes & surcharges"}
FIX = ["Key", "Entity", "Sector", "Sub-segment", "Role", "Currency", "Listing / ownership", "LC per USD"]
C25 = {f: 9 + i for i, f in enumerate(FIELDS)}
C24 = {f: 9 + len(FIELDS) + i for i, f in enumerate(FIELDS)}
NOTE_COL = 9 + 2 * len(FIELDS)
SRC_COL = NOTE_COL + 1
ds.cell(row=3, column=9, value="FY2025 (year to 31-Dec-2025)").font = f_bold
ds.cell(row=3, column=9 + len(FIELDS), value="FY2024 (year to 31-Dec-2024)").font = f_bold
heads = FIX + [LABEL[f] for f in FIELDS] * 2 + ["Notes", "Source"]
for j, h in enumerate(heads, 1):
    c = ds.cell(row=4, column=j, value=h)
    c.font, c.alignment, c.border = f_hdr, ctr, box
    c.fill = fill_hdr if j < 9 or j == NOTE_COL or j == SRC_COL else (fill_grp if j < 9 + len(FIELDS) else fill_hdr)
ds.row_dimensions[4].height = 42
ROW = {}
FORMULA_FIELDS = {"gross_profit", "ebitda"}
for r, e in enumerate(D["entities"], 5):
    ROW[e["key"]] = r
    for j, v in enumerate([e["key"], e["name"], e["sector"], e["sub"], e["role"], e["ccy"], e["listing"]], 1):
        ds.cell(row=r, column=j, value=v).border = box
    c = ds.cell(row=r, column=8, value=f'=VLOOKUP(F{r},FX!$A$5:$B$8,2,FALSE)')
    c.number_format, c.border, c.font = "#,##0.000", box, f_calc
    for yr, cols in (("2025", C25), ("2024", C24)):
        row = e["fy"][yr]
        for f in FIELDS:
            cell = ds.cell(row=r, column=cols[f])
            cell.border = box
            if f == "gross_profit":
                cell.value = f'=IF({L(cols["cogs"])}{r}="","",{L(cols["revenue"])}{r}-{L(cols["cogs"])}{r})'
                cell.font = f_calc
            elif f == "ebitda":
                cell.value = f'={L(cols["ebit"])}{r}+{L(cols["da"])}{r}'
                cell.font = f_calc
            else:
                v = row.get(f)
                cell.value = None if v is None else round(v, 3)
                cell.font = f_in
            cell.number_format = '#,##0.0;[Red]-#,##0.0'
    ds.cell(row=r, column=NOTE_COL, value=e["note"]).alignment = left
    ds.cell(row=r, column=SRC_COL, value=e["source"]).alignment = left
    if e["role"] == "Sinochem":
        for j in range(1, 9):
            ds.cell(row=r, column=j).fill = fill_sin
ds.freeze_panes = "C5"
for j, w in {1: 6, 2: 34, 3: 12, 4: 30, 5: 10, 6: 9, 7: 26, 8: 11}.items():
    ds.column_dimensions[L(j)].width = w
for j in range(9, NOTE_COL):
    ds.column_dimensions[L(j)].width = 13
ds.column_dimensions[L(NOTE_COL)].width, ds.column_dimensions[L(SRC_COL)].width = 60, 44

# ----------------------------------------------------------------------------------------------------------------------- metric catalogue
def dref(key, field, yr="2025"):
    col = C25[field] if yr == "2025" else C24[field]
    return f"Data!${L(col)}${ROW[key]}"


def metric_formulas(k):
    c = lambda f: dref(k, f)          # noqa: E731
    p = lambda f: dref(k, f, "2024")  # noqa: E731
    fxr = f"Data!$H${ROW[k]}"
    avg = lambda f: f"AVERAGE({c(f)},{p(f)})"  # noqa: E731
    return {
        "rev_usd": f"={c('revenue')}/{fxr}/1000",
        "ta_usd": f"={c('total_assets')}/{fxr}/1000",
        "g_rev": f'=IFERROR({c("revenue")}/{p("revenue")}-1,"n/a")',
        "g_ebitda": f'=IFERROR(IF({p("ebitda")}<=0,"n.m.",{c("ebitda")}/{p("ebitda")}-1),"n/a")',
        "gm": f'=IF({c("gross_profit")}="","n/a",{c("gross_profit")}/{c("revenue")})',
        "gm_tax": f'=IF(OR({c("gross_profit")}="",{c("tax_surcharge")}=""),"n/a",({c("gross_profit")}-{c("tax_surcharge")})/{c("revenue")})',
        "m_ebitda": f'=IFERROR({c("ebitda")}/{c("revenue")},"n/a")',
        "m_ebit": f'=IFERROR({c("ebit")}/{c("revenue")},"n/a")',
        "m_net": f'=IFERROR({c("net_income")}/{c("revenue")},"n/a")',
        "ebitda_assets": f'=IFERROR({c("ebitda")}/{avg("total_assets")},"n/a")',
        "roe": f'=IFERROR({c("net_income")}/{avg("total_equity")},"n/a")',
        "roa": f'=IFERROR({c("net_income")}/{avg("total_assets")},"n/a")',
        "turn": f'=IFERROR({c("revenue")}/{avg("total_assets")},"n/a")',
        "liab": f'=IFERROR({c("total_liabilities")}/{c("total_assets")},"n/a")',
        "de": f'=IFERROR({c("debt")}/{c("total_equity")},"n/a")',
        "nd_ebitda": f'=IFERROR(IF({c("ebitda")}<=0,"n.m.",({c("debt")}-{c("cash")})/{c("ebitda")}),"n/a")',
        "icov": f'=IFERROR(IF({c("interest_exp")}<=0,"n/a",{c("ebit")}/{c("interest_exp")}),"n/a")',
        "cur": f'=IFERROR({c("current_assets")}/{c("current_liabilities")},"n/a")',
        "ocf_rev": f'=IFERROR({c("ocf")}/{c("revenue")},"n/a")',
        "ocf_ebitda": f'=IFERROR(IF({c("ebitda")}<=0,"n.m.",{c("ocf")}/{c("ebitda")}),"n/a")',
        "capex_rev": f'=IFERROR({c("capex")}/{c("revenue")},"n/a")',
        "fcf_m": f'=IFERROR(({c("ocf")}-{c("capex")})/{c("revenue")},"n/a")',
        "fcf_usd": f"=({c('ocf')}-{c('capex')})/{fxr}/1000",
    }


# key, header, group, number format, direction (1 higher better, -1 lower better, 0 informational)
METRICS = [
    ("rev_usd", "Revenue (USD bn)", "Scale", '#,##0.0', 0), ("ta_usd", "Total assets (USD bn)", "Scale", '#,##0.0', 0),
    ("g_rev", "Revenue growth FY25", "Growth", '0.0%', 1), ("g_ebitda", "EBITDA growth FY25", "Growth", '0.0%', 1),
    ("gm", "Gross margin", "Profitability", '0.0%', 1), ("gm_tax", "Gross margin after taxes & surcharges", "Profitability", '0.0%', 1),
    ("m_ebitda", "EBITDA margin", "Profitability", '0.0%', 1), ("m_ebit", "EBIT margin", "Profitability", '0.0%', 1),
    ("m_net", "Net margin", "Profitability", '0.0%', 1), ("ebitda_assets", "EBITDA / avg assets", "Returns", '0.0%', 1),
    ("roe", "ROE (NI / avg equity)", "Returns", '0.0%', 1), ("roa", "ROA (NI / avg assets)", "Returns", '0.0%', 1),
    ("turn", "Asset turnover (x)", "Returns", '0.00"x"', 1),
    ("liab", "Liabilities / assets", "Leverage & liquidity", '0.0%', -1), ("de", "Debt / equity", "Leverage & liquidity", '0.0%', -1),
    ("nd_ebitda", "Net debt / EBITDA (x)", "Leverage & liquidity", '0.00"x"', -1), ("icov", "Interest cover (EBIT / interest, x)", "Leverage & liquidity", '0.0"x"', 1),
    ("cur", "Current ratio (x)", "Leverage & liquidity", '0.00"x"', 1),
    ("ocf_rev", "OCF / revenue", "Cash generation", '0.0%', 1), ("ocf_ebitda", "OCF / EBITDA (cash conversion)", "Cash generation", '0.0%', 1),
    ("capex_rev", "Capex / revenue", "Cash generation", '0.0%', 0), ("fcf_m", "FCF margin", "Cash generation", '0.0%', 1),
    ("fcf_usd", "Free cash flow (USD bn)", "Cash generation", '#,##0.00;[Red]-#,##0.00', 1),
]
MCOL = {m[0]: 4 + i for i, m in enumerate(METRICS)}   # first metric in column D

SECTORS = {
    "Energy": dict(
        title="Energy - Sinochem Energy vs listed refining / petrochemical peers",
        sub="Sinochem Energy (中化能源) is the operating arm of Sinochem Corp's petrochemical segment: crude/product trading, 15 mt/y of refining capacity centred on Quanzhou, and fuel retail. FY2025 audited statements.",
        sin=[("SCE", "Sinochem")],
        peers=[("SNP", "State-owned major"), ("PTR", "State-owned major"), ("HLP", "Private integrated refiner"), ("RSP", "Private integrated refiner"),
               ("SSP", "Refining & petchem pure-play")],
        subs=[("State-owned majors", ["SNP", "PTR"]), ("Private integrated refiners", ["HLP", "RSP"])]),
    "Agriculture": dict(
        title="Agriculture - Syngenta Group entities vs crop-protection / seeds and crop-nutrition peers",
        sub="Syngenta Group itself discloses only sales and EBITDA publicly (USD 28.4bn / USD 4.4bn in FY2025); full statements exist for Syngenta AG (global CP + Seeds), ADAMA, Yangnong and Sinofert.",
        sin=[("SYN", "Sinochem - CP & seeds"), ("ADM", "Sinochem - CP (generics)"), ("YNG", "Sinochem - CP (AI)"), ("SIF", "Sinochem - fertilizer")],
        peers=[("CTV", "CP & seeds peer"), ("FMC", "CP & seeds peer"), ("JSA", "CP & seeds peer"), ("LCH", "CP & seeds peer"),
               ("NTR", "Crop-nutrition peer"), ("MOS", "Crop-nutrition peer"), ("YTH", "Crop-nutrition peer"), ("SLP", "Crop-nutrition peer")],
        subs=[("Crop-protection & seeds peers", ["CTV", "FMC", "JSA", "LCH"]), ("Crop-nutrition peers", ["NTR", "MOS", "YTH", "SLP"])]),
    "Rubber": dict(
        title="Rubber & tyres - Pirelli / Aeolus vs global and Chinese tyre peers and a natural-rubber producer",
        sub="Sinochem's rubber & tyre exposure is held through the ChemChina line (China National Tire & Rubber): Pirelli (Sinochem holds ~34.1% via Marco Polo International Italy; Pirelli's accounts state no entity controls it) and Aeolus Tyre.",
        sin=[("PIR", "Sinochem - premium tyres"), ("AEO", "Sinochem - truck & bus tyres")],
        peers=[("MIC", "Global major"), ("BRI", "Global major"), ("GYR", "Global major"), ("LLT", "China tyre producer"), ("SLG", "China tyre producer"),
               ("ZCR", "China tyre producer"), ("HNR", "Natural rubber")],
        subs=[("Global majors", ["MIC", "BRI", "GYR"]), ("China tyre producers", ["LLT", "SLG", "ZCR"])]),
}

LAYOUT = {}  # sector -> dict(rows..., stats rows) used by the scorecard


def build_sector(name, cfg):
    ms = [m for m in METRICS if not (m[0] == "gm_tax" and name != "Energy")]
    mcol = {m[0]: 4 + i for i, m in enumerate(ms)}
    ws = wb.create_sheet(name)
    ws["A1"], ws["A1"].font = cfg["title"], f_title
    ws["A2"], ws["A2"].font = cfg["sub"], f_sub
    ws["A3"] = "FY2025, ratios computed from the 'Data' sheet. Amber rows = Sinochem Holdings group entities. Peer statistics exclude Sinochem entities. 'n.m.' = not meaningful (negative EBITDA); 'n/a' = not available."
    ws["A3"].font = f_sub
    hdr_r, dir_r = 6, 7
    for j, h in enumerate(["Entity", "Group", "Ccy"], 1):
        ws.cell(row=hdr_r, column=j, value=h)
    for key, h, grp, fmt, d in ms:
        ws.cell(row=hdr_r, column=mcol[key], value=h)
        g = ws.cell(row=hdr_r - 1, column=mcol[key], value=grp)
        g.font, g.fill, g.alignment = f_hdr, fill_grp, ctr
        dc = ws.cell(row=dir_r, column=mcol[key], value=d)
        dc.font, dc.alignment, dc.number_format = f_sub, ctr, '"higher is better";"lower is better";"informational"'
    ws.cell(row=dir_r, column=1, value="Direction (+1 higher is better, -1 lower is better, 0 informational)").font = f_sub
    for j in range(1, 4 + len(ms)):
        c = ws.cell(row=hdr_r, column=j)
        c.font, c.fill, c.alignment, c.border = f_hdr, fill_hdr, ctr, box
    ws.row_dimensions[hdr_r].height = 56
    start = None
    cols = [mcol[m[0]] for m in ms]
    groups = [m[2] for m in ms]
    i = 0
    while i < len(ms):
        j = i
        while j + 1 < len(ms) and groups[j + 1] == groups[i]:
            j += 1
        if j > i:
            ws.merge_cells(start_row=hdr_r - 1, start_column=cols[i], end_row=hdr_r - 1, end_column=cols[j])
        i = j + 1
    r = dir_r + 1
    rows = {}
    for grp_name, members in (("sin", cfg["sin"]), ("peer", cfg["peers"])):
        for key, label in members:
            e = ENT[key]
            rows[key] = r
            ws.cell(row=r, column=1, value=e["name"]).font = f_bold if grp_name == "sin" else f_calc
            ws.cell(row=r, column=2, value=label).font = f_calc
            ws.cell(row=r, column=3, value=e["ccy"]).font = f_calc
            fm = metric_formulas(key)
            for mk, h, g, fmt, d in ms:
                c = ws.cell(row=r, column=mcol[mk], value=fm[mk])
                c.number_format, c.alignment, c.font = fmt, Alignment(horizontal="right"), f_calc
            for j in range(1, 4 + len(ms)):
                cell = ws.cell(row=r, column=j)
                cell.border = box
                if grp_name == "sin":
                    cell.fill = fill_sin
            r += 1
    first_peer, last_peer = rows[cfg["peers"][0][0]], rows[cfg["peers"][-1][0]]
    r += 1
    stats = {}

    def stat_row(label, fn, rng_fn, style=f_bold):
        nonlocal r
        ws.cell(row=r, column=1, value=label).font = style
        for mk, h, g, fmt, d in ms:
            col = L(mcol[mk])
            c = ws.cell(row=r, column=mcol[mk], value=f'=IFERROR({fn}({rng_fn(col)}),"n/a")')
            c.number_format, c.alignment, c.font = fmt, Alignment(horizontal="right"), style
        for j in range(1, 4 + len(ms)):
            ws.cell(row=r, column=j).fill, ws.cell(row=r, column=j).border = fill_stat, box
        stats[label] = r
        r += 1

    allr = lambda col: f"{col}{first_peer}:{col}{last_peer}"  # noqa: E731
    stat_row("Peer median - all peers", "MEDIAN", allr)
    stat_row("Peer high - all peers", "MAX", allr, f_calc)
    stat_row("Peer low - all peers", "MIN", allr, f_calc)
    sub_rows = {}
    for sub_name, members in cfg["subs"]:
        a, b = rows[members[0]], rows[members[-1]]
        stat_row(f"Peer median - {sub_name}", "MEDIAN", lambda col, a=a, b=b: f"{col}{a}:{col}{b}")
        sub_rows[sub_name] = (a, b)
    # anchor vs median
    anchor = cfg["sin"][0][0]
    r += 1
    ws.cell(row=r, column=1, value=f"Anchor ({ENT[anchor]['name']}) minus all-peer median").font = f_bold
    delta_r = r
    for mk, h, g, fmt, d in ms:
        col = L(mcol[mk])
        c = ws.cell(row=r, column=mcol[mk], value=f'=IFERROR({col}{rows[anchor]}-{col}{stats["Peer median - all peers"]},"n/a")')
        c.number_format, c.alignment, c.font = fmt.replace('0.0%', '0.0%;-0.0%'), Alignment(horizontal="right"), f_bold
    for j in range(1, 4 + len(ms)):
        ws.cell(row=r, column=j).border = box
    # conditional colours on the delta row: green = better than median, red = worse
    for mk, h, g, fmt, d in ms:
        col = L(mcol[mk])
        rng = f"{col}{delta_r}"
        if d != 0:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rng}),{rng}*{col}${dir_r}>0)'], font=Font(color="1B7F3B", bold=True)))
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rng}),{rng}*{col}${dir_r}<0)'], font=Font(color="C0392B", bold=True)))
    ws.freeze_panes = ws.cell(row=dir_r + 1, column=4)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 40, 26, 6
    for mk, *_ in ms:
        ws.column_dimensions[L(mcol[mk])].width = 13.5
    LAYOUT[name] = dict(mcol=mcol, rows=rows, stats=stats, first_peer=first_peer, last_peer=last_peer, subs=sub_rows, cfg=cfg)
    return ws


for nm, cfg in SECTORS.items():
    build_sector(nm, cfg)

# ----------------------------------------------------------------------------------------------------------------------- Scorecard
sc = wb.create_sheet("Scorecard", 0)
sc["A1"], sc["A1"].font = "Scorecard - where each Sinochem entity ranks against its peer group (FY2025)", f_title
sc["A2"] = ("Rank 1 = best among the entity plus its peers (n shown). Tier: top third / mid / bottom third by rank. "
            "All values are live formulas pulling from the sector sheets. Syngenta AG, ADAMA and Yangnong are compared with crop-protection & seeds peers; Sinofert with crop-nutrition peers.")
sc["A2"].font = f_sub
SC_METRICS = ["g_rev", "m_ebitda", "m_net", "roe", "ebitda_assets", "nd_ebitda", "icov", "ocf_ebitda", "fcf_m"]
SC_PEERS = {  # entity -> (sector, peer range builder)
    "SCE": ("Energy", None), "SYN": ("Agriculture", "Crop-protection & seeds peers"), "ADM": ("Agriculture", "Crop-protection & seeds peers"),
    "YNG": ("Agriculture", "Crop-protection & seeds peers"), "SIF": ("Agriculture", "Crop-nutrition peers"), "PIR": ("Rubber", None), "AEO": ("Rubber", None),
}
hdr = ["Sinochem entity", "Peer group", "Metric", "Entity", "Peer median", "Entity - median", "Rank", "of n", "Tier"]
for j, h in enumerate(hdr, 1):
    c = sc.cell(row=4, column=j, value=h)
    c.font, c.fill, c.alignment, c.border = f_hdr, fill_hdr, ctr, box
r = 5
name_of = {m[0]: m[1] for m in METRICS}
dir_of = {m[0]: m[4] for m in METRICS}
fmt_of = {m[0]: m[3] for m in METRICS}
for key, (sector, sub) in SC_PEERS.items():
    lay = LAYOUT[sector]
    a, b = (lay["first_peer"], lay["last_peer"]) if sub is None else lay["subs"][sub]
    peer_label = "All listed peers" if sub is None else sub
    ent_row = lay["rows"][key]
    start = r
    for mk in SC_METRICS:
        col = L(lay["mcol"][mk])
        rng = f"'{sector}'!{col}{a}:{col}{b}"
        ent = f"'{sector}'!{col}{ent_row}"
        d = dir_of[mk]
        sc.cell(row=r, column=1, value=ENT[key]["name"])
        sc.cell(row=r, column=2, value=peer_label)
        sc.cell(row=r, column=3, value=name_of[mk])
        sc.cell(row=r, column=4, value=f"={ent}").number_format = fmt_of[mk]
        sc.cell(row=r, column=5, value=f'=IFERROR(MEDIAN({rng}),"n/a")').number_format = fmt_of[mk]
        sc.cell(row=r, column=6, value=f'=IFERROR(D{r}-E{r},"n/a")').number_format = fmt_of[mk].replace('0.0%', '0.0%;-0.0%')
        better = ">" if d == 1 else "<"
        sc.cell(row=r, column=7, value=f'=IF(ISNUMBER(D{r}),1+COUNTIF({rng},"{better}"&D{r}),"n/a")')
        sc.cell(row=r, column=8, value=f'=IF(ISNUMBER(D{r}),1+COUNT({rng}),"n/a")')
        sc.cell(row=r, column=9, value=f'=IF(ISNUMBER(G{r}),IF((H{r}-G{r})/(H{r}-1)>=0.67,"Top third",IF((H{r}-G{r})/(H{r}-1)<=0.33,"Bottom third","Mid-pack")),'
                                       f'IF(D{r}="n.m.","Bottom third (n.m.)","n/a"))')
        for j in range(1, 10):
            cell = sc.cell(row=r, column=j)
            cell.border = box
            cell.font = f_bold if j == 1 else f_calc
            if j in (4, 5, 6):
                cell.alignment = Alignment(horizontal="right")
            if j in (7, 8, 9):
                cell.alignment = Alignment(horizontal="center")
        r += 1
    sc.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    sc.merge_cells(start_row=start, start_column=2, end_row=r - 1, end_column=2)
    for rr in range(start, r):
        for j in (1, 2):
            sc.cell(row=rr, column=j).fill = fill_sin
            sc.cell(row=rr, column=j).alignment = Alignment(vertical="center", wrap_text=True)
rng = f"I5:I{r - 1}"
sc.conditional_formatting.add(rng, FormulaRule(formula=['I5="Top third"'], font=Font(color="1B7F3B", bold=True), fill=PatternFill(start_color="DFF0E3", end_color="DFF0E3", fill_type="solid")))
sc.conditional_formatting.add(rng, FormulaRule(formula=['LEFT(I5,12)="Bottom third"'], font=Font(color="C0392B", bold=True), fill=PatternFill(start_color="F8DCD8", end_color="F8DCD8", fill_type="solid")))
sc.freeze_panes = "A5"
for col, w in zip("ABCDEFGHI", [32, 28, 32, 12, 12, 14, 7, 7, 18]):
    sc.column_dimensions[col].width = w

# ----------------------------------------------------------------------------------------------------------------------- Interim
it = wb.create_sheet("H1-2026 momentum")
it["A1"], it["A1"].font = "Latest interim read-through: H1 2026 vs H1 2025 (cumulative six months, reporting currency, millions)", f_title
it["A2"] = ("Available for China A-share companies and Sinofert from their H1 2026 reports (published Aug 2026). Sinochem Energy, Syngenta AG and Pirelli interim items are in 'Group context'.")
it["A2"].font = f_sub
hh = ["Entity", "Sector", "Role", "Revenue H1-26", "Revenue H1-25", "Revenue YoY", "Net income H1-26", "Net income H1-25", "Net income change", "Net margin H1-26", "Net margin H1-25"]
for j, h in enumerate(hh, 1):
    c = it.cell(row=4, column=j, value=h)
    c.font, c.fill, c.alignment, c.border = f_hdr, fill_hdr, ctr, box
r = 5
order = ["SCE", "SNP", "PTR", "HLP", "RSP", "SSP", "ADM", "YNG", "SIF", "JSA", "LCH", "YTH", "SLP", "AEO", "LLT", "SLG", "ZCR", "HNR"]
for k in order:
    if k not in D["interim"]:
        continue
    e, v = ENT[k], D["interim"][k]
    it.cell(row=r, column=1, value=e["name"])
    it.cell(row=r, column=2, value=e["sector"])
    it.cell(row=r, column=3, value=e["role"])
    for j, val in ((4, v["h1_2026"]["revenue"]), (5, v["h1_2025"]["revenue"]), (7, v["h1_2026"]["net_income"]), (8, v["h1_2025"]["net_income"])):
        c = it.cell(row=r, column=j, value=round(val, 3))
        c.font, c.number_format = f_in, '#,##0;[Red]-#,##0'
    it.cell(row=r, column=6, value=f'=IFERROR(D{r}/E{r}-1,"n/a")').number_format = '0.0%;-0.0%'
    it.cell(row=r, column=9, value=f'=G{r}-H{r}').number_format = '#,##0;[Red]-#,##0'
    it.cell(row=r, column=10, value=f'=IFERROR(G{r}/D{r},"n/a")').number_format = '0.0%;-0.0%'
    it.cell(row=r, column=11, value=f'=IFERROR(H{r}/E{r},"n/a")').number_format = '0.0%;-0.0%'
    for j in range(1, 12):
        it.cell(row=r, column=j).border = box
        if e["role"] == "Sinochem" and j < 4:
            it.cell(row=r, column=j).fill = fill_sin
    r += 1
it.freeze_panes = "B5"
for col, w in zip("ABCDEFGHIJK", [36, 12, 10, 14, 14, 12, 16, 16, 16, 14, 14]):
    it.column_dimensions[col].width = w
it.cell(row=r + 1, column=1, value="Source: Eastmoney F10 (H1 2026 and H1 2025 cumulative income statements as filed). Sinochem Energy has not published an H1 2026 statement that I could locate; see 'Group context' for Sinochem Corp Q1 2026.").font = f_sub

# ----------------------------------------------------------------------------------------------------------------------- Group context
gc = wb.create_sheet("Group context")
gc["A1"], gc["A1"].font = "Group context - what is disclosed at Sinochem level", f_title
gc["A2"] = ("Sinochem Holdings (中国中化控股) is unlisted and does not publish a consolidated sector split. The tables below are the disclosed building blocks. RMB 100m (亿元) unless stated.")
gc["A2"].font = f_sub


def put_table(r0, title, header, rows, fmt='#,##0.00;[Red]-#,##0.00', src=None):
    gc.cell(row=r0, column=1, value=title).font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for j, h in enumerate(header, 1):
        c = gc.cell(row=r0 + 1, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = f_hdr, fill_hdr, ctr, box
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            c = gc.cell(row=r0 + 2 + i, column=j, value=v)
            c.border = box
            if isinstance(v, (int, float)):
                c.font, c.number_format = f_in, fmt
    end = r0 + 2 + len(rows)
    if src:
        gc.cell(row=end, column=1, value=src).font = f_sub
        end += 1
    return end + 1


r = 4
r = put_table(r, "1. 中国中化股份有限公司 (Sinochem Corp, main financing platform of Sinochem Holdings) - consolidated", ["Item (RMB 100m)", "FY2023", "FY2024", "FY2025", "Q1 2026"], [
    ["Revenue", 6170.73, 6083.80, 5649.67, 1195.77], ["Gross margin (%)", 4.93, 4.72, 4.61, 5.42], ["Operating business profit", 31.68, 32.88, 13.62, 5.41],
    ["Asset impairment loss", -70.59, -46.16, -31.45, -0.95], ["Investment income", 54.12, 121.87, 73.13, 2.00], ["Fair-value change gain/(loss)", -4.93, -18.42, 42.47, 27.55],
    ["Total profit", 3.50, 86.59, 86.68, 42.19], ["Total assets", 8565.46, 9187.36, 9685.50, 10241.52], ["Total debt", 4110.07, 4632.81, 4839.40, 5656.67],
    ["Total liabilities", 5781.43, 6078.22, 6120.83, 6492.67], ["Owner's equity (incl. RMB 98bn perpetuals at 1Q25)", 2784.03, 3109.14, 3564.67, 3748.85],
    ["Total capitalisation ratio (%)", 66.56, 68.56, 68.21, 69.87], ["Operating cash flow", 269.02, 92.57, 193.67, -172.41], ["Total debt / EBITDA (x)", 20.23, 14.12, 15.28, None],
    ["EBITDA interest cover (x)", 1.79, 2.44, 2.52, None]],
    src="Source: 中诚信国际 2026 tracking rating report for 中国中化股份有限公司 (SSE, 2026-06-23), table 7; 2025 report for FY2023 comparatives.")
r = put_table(r, "2. Sinochem Corp segment revenue and gross margin (RMB 100m, %)", ["Segment", "Rev FY2023", "GM% FY2023", "Rev FY2024", "GM% FY2024", "Rev FY2025", "GM% FY2025"], [
    ["Petrochemicals (石油化工)", 4573.07, 2.90, 4300.80, 2.17, 3959.51, 1.81], ["Materials science (材料科学)", 547.00, 4.88, 817.95, 6.86, 745.79, 7.65],
    ["Urban operations (城市运营)", 716.66, 12.22, 590.47, 14.57, 593.71, 15.86]],
    src="Source: same rating reports, table 6. Agriculture and rubber/tyre sit in the ChemChina line, outside Sinochem Corp's consolidation.")
r = put_table(r, "3. Sinochem Energy (中化能源) revenue by segment (RMB 100m)", ["Segment", "FY2023", "FY2024", "FY2025"], [
    ["Oil trading & services", 3418.87, 3048.98, 2736.92], ["Oil products & chemicals sales", 1796.10, 1725.78, 1606.52], ["Refining & chemicals", 822.12, 734.93, 612.66],
    ["Storage & logistics", 17.70, 17.76, 24.20], ["Other", 11.24, 8.10, 7.61], ["Inter-segment elimination", -1483.96, -1266.53, -1046.80], ["Total revenue", 4582.06, 4269.03, 3941.10],
    ["Net profit/(loss)", None, -31.82, -16.83]],
    src="Source: rating report table 1 (Sinochem Energy annual reports); FY2025 net loss ties to audited statements (RMB -1,682.7m).")
r = put_table(r, "4. Syngenta Group headline (USD bn) - consolidated, unaudited", ["Item", "FY2024", "FY2025", "H1 2025", "H1 2026"], [
    ["Sales", 28.8, 28.4, None, 12.2], ["EBITDA", 3.9, 4.4, None, 2.4], ["EBITDA margin (%)", 13.5, 15.4, 18.6, 19.5],
    ["Syngenta Crop Protection sales", 13.2, 13.7, None, 6.6], ["Seeds sales", 4.8, 4.8, None, None], ["ADAMA sales", 4.1, 4.1, None, None],
    ["Syngenta Group China sales", 9.2, 8.3, None, None]], fmt='#,##0.0',
    src="Source: Syngenta Group FY2025 results (2026-03-31) and H1 2026 results (2026-08-27). Sinofert was transferred to Sinochem Hong Kong on 23 Dec 2025 and is not consolidated from Jan 2026; H1 2026 is like-for-like.")
r = put_table(r, "5. Pirelli - latest data points (EUR m)", ["Item", "FY2024", "FY2025", "H1 2025", "H1 2026"], [
    ["Revenue", 6773.3, 6776.2, 3498.6, 3494.5], ["Adjusted EBIT", 1060.5, 1081.4, None, 557.8], ["Net income", 501.1, 530.7, None, 299.0],
    ["Net financial position (debt)", 1925.8, 1102.0, None, 1915.9]], fmt='#,##0.0',
    src="Source: Pirelli FY2025 and H1 2026 results releases. H1 2026 net income was +13.3% year on year. Directors Zhang Haitao, Xi Xiaohong and Wang Kun voted against the H1 2026 report over its declaration of control by MTP S.p.A. Sinochem (via MPI Italy) held ~34.1% at Dec 2025 (AR 2025).")
gc.column_dimensions["A"].width = 48
for col in "BCDEFG":
    gc.column_dimensions[col].width = 15

# ----------------------------------------------------------------------------------------------------------------------- README
rd = wb.create_sheet("README", 0)
lines = [
    ("Sinochem Holdings - sector peer comparison (Energy, Agriculture, Rubber)", f_title),
    ("Prepared 12 Oct 2026. Basis: fiscal year ended 31 Dec 2025 (audited / filed), with H1 2026 momentum where available.", f_sub),
    ("", None),
    ("How to read this workbook", f_bold),
    ("Scorecard - rank of each Sinochem entity against its peer group on nine headline metrics (live formulas).", None),
    ("Energy / Agriculture / Rubber - full ratio matrices, peer median/high/low, sub-group medians, and anchor-minus-median row (green = better, red = worse).", None),
    ("H1-2026 momentum - first-half 2026 vs first-half 2025 revenue and net income.", None),
    ("Group context - disclosed Sinochem-level data (Sinochem Corp, Sinochem Energy segments, Syngenta Group, Pirelli).", None),
    ("Data - standardised inputs (blue) with the FY2024 comparatives used for growth and averages; FX - editable USD translation rates.", None),
    ("", None),
    ("What 'Sinochem' means here", f_bold),
    ("Sinochem Holdings (中国中化控股) is a state-owned, unlisted group; it publishes no consolidated sector split. Each sector is therefore anchored on the group's operating entities that DO disclose full financials:", None),
    ("  Energy: Sinochem Energy (中化能源), the petrochemical segment's operating company (trading + 15 mt/y refining capacity centred on Quanzhou + retail).", None),
    ("  Agriculture: Syngenta AG (global crop protection + seeds) with ADAMA, Yangnong and Sinofert as additional group entities; Syngenta Group headline sales/EBITDA in 'Group context'.", None),
    ("  Rubber: Pirelli (Sinochem holds ~34.1%; Pirelli's accounts say no entity controls it, so it is not consolidated by Sinochem) and Aeolus Tyre. Sinochem International agreed in Nov 2022 to sell 36% of Halcyon Agri (natural rubber) to Hainan Rubber (completion not verified here).", None),
    ("", None),
    ("Definitions (applied identically to every company)", f_bold),
    ("EBIT = operating profit before financing as reported (China GAAP: operating profit + net finance expense; IFRS/US GAAP: reported operating income or closest equivalent; Corteva: pre-tax profit + interest expense).", None),
    ("EBITDA = EBIT + depreciation & amortisation; reported basis, NOT adjusted for impairments, restructuring or litigation (so it will differ from company 'adjusted EBITDA').", None),
    ("Debt = bank borrowings + bonds + lease liabilities (US GAAP filers exclude operating leases). Net debt = Debt - cash (no credit for deposits or other financial assets).", None),
    ("ROE / ROA use net income (incl. minorities) over the average of opening and closing total equity / assets. Capex = cash paid for PP&E and intangibles.", None),
    ("Gross margin after taxes & surcharges deducts China's consumption tax and levies, which Chinese refiners report below the gross-profit line - needed to compare Sinochem Energy with Chinese peers.", None),
    ("", None),
    ("Key caveats", f_bold),
    ("1) Entity-level, not group-level: Sinochem Holdings' consolidated sector results are not disclosed; unlisted entities (Sinochem Energy, Syngenta AG) are compared with listed peers on filed statements.", None),
    ("2) Business-mix differences matter: crude/product trading (55%) and fuel/chemical wholesale-retail (32%) make up ~87% of Sinochem Energy's pre-elimination segment revenue (refining 12%), so gross margin is structurally thin; read EBITDA/assets and cash metrics for like-for-like comparison.", None),
    ("3) Reported-basis EBITDA includes one-offs: Syngenta AG (USD 792m litigation + USD 225m restructuring), FMC and Goodyear (large 2025 charges), Sinochem Energy (RMB 1.0bn fair-value gain).", None),
    ("4) Different fiscal periods/standards: all peers are Dec-2025 year-ends; GAAP differs (China GAAP / IFRS / US GAAP). Pirelli has no cost-of-sales line (by-nature P&L).", None),
    ("5) USD translation uses 2025 average rates for both P&L and balance-sheet items (indicative sizes only; ratios are currency-neutral).", None),
    ("6) Interim read-through (H1 2026) is limited to companies that filed H1 reports; the 2026 oil-price spike materially lifts refining profits.", None),
]
for i, (t, f) in enumerate(lines, 1):
    c = rd.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if f:
        c.font = f
    rd.row_dimensions[i].height = 15 * max(1, -(-len(t) // 135))
rd.column_dimensions["A"].width = 130

# ----------------------------------------------------------------------------------------------------------------------- Sources
so = wb.create_sheet("Sources")
so["A1"], so["A1"].font = "Sources", f_title
src_rows = [
    ("Sinochem Corp 2026 tracking rating report (FY2025, Q1 2026)", "http://static.sse.com.cn/disclosure/bond/announcement/company/c/new/2026-06-23/244212_20260623_CREQ.pdf"),
    ("Sinochem Corp 2025 tracking rating report (FY2024)", "http://static.sse.com.cn/disclosure/bond/announcement/company/c/new/2025-06-23/241247_20250623_1DU5.pdf"),
    ("Sinochem Energy FY2025 audited consolidated statements (天职业字[2026]19618号)", "http://static.sse.com.cn/disclosure/bond/announcement/company/c/new/2026-04-30/115421_20260430_T7YL.pdf"),
    ("Syngenta AG Financial Report 2025", "https://www.syngenta.com/sites/default/files/2026-03/Financial%20Report%202025_0.pdf"),
    ("Syngenta Group FY2025 results release", "https://www.syngentagroup.com/newsroom/2026/syngenta-group-reports-2025-full-year-results"),
    ("Syngenta Group H1 2026 results", "https://www.syngentagroup.com/sites/default/files/2026-08/260827-Syngenta-Group-Q2-H1_EN_2949a0aa.pdf"),
    ("Pirelli 2025 Annual Report", "https://corp-assets.pirelli.com/corporate/Annual_Report_Pirelli_2025_ENG.pdf"),
    ("Pirelli FY2025 preliminary results release", "https://press.pirelli.com/review-of-preliminary-results-to-31-december-2025-and-2026-budget/"),
    ("Pirelli H1 2026 results release", "https://press.pirelli.com/pirelli-board-majority-approves-consolidated-results-to-30-june-2026/"),
    ("Michelin 2025 Universal Registration Document", "https://www.michelin.com/en/publications/group/financial-information-december-31-2025"),
    ("Bridgestone Consolidated Financial Results FY2025", "https://www.bridgestone.com/ir/library/result/pdf/e_r7_4_results.pdf"),
    ("Nutrien 2025 audited financial statements (40-F ex. 99.3)", "https://www.sec.gov/Archives/edgar/data/1725964/000119312526081326/d56746dex993.htm"),
    ("SEC EDGAR XBRL company facts (Corteva, FMC, Mosaic, Goodyear)", "https://data.sec.gov/api/xbrl/companyfacts/"),
    ("Eastmoney F10 financial statements (A-shares, Sinofert, US balance/income)", "https://datacenter.eastmoney.com/securities/api/data/v1/get"),
    ("IRS yearly average exchange rates", "https://www.irs.gov/individuals/international-taxpayers/yearly-average-currency-exchange-rates"),
    ("Sinochem International sale of 36% of Halcyon Agri to Hainan Rubber (Nov 2022)", "https://www.stcn.com/article/detail/733704.html"),
    ("Michelin H1 2026 results", "https://www.michelin.com/en/publications/group/financial-information-at-june-30-2026"),
]
for i, (n, u) in enumerate(src_rows, 3):
    so.cell(row=i, column=1, value=n)
    so.cell(row=i, column=2, value=u)
so.column_dimensions["A"].width, so.column_dimensions["B"].width = 80, 140

for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"
    _ws.page_setup.fitToWidth = 1
    _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr.fitToPage = True
    _ws.sheet_view.showGridLines = False

wb._sheets = [wb[n] for n in ["README", "Scorecard", "Energy", "Agriculture", "Rubber", "H1-2026 momentum", "Group context", "Data", "FX", "Sources"]]
out = ROOT / "Sinochem_Peer_Comparison.xlsx"
wb.save(out)
print("saved", out)
json.dump({k: {"rows": v["rows"], "stats": v["stats"], "first_peer": v["first_peer"], "last_peer": v["last_peer"]} for k, v in LAYOUT.items()},
          open(ROOT / "data" / "layout.json", "w"))
