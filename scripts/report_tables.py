"""Render the core of each sector matrix as markdown, straight from the shipped workbook's computed values."""
import json, pathlib, sys
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(ROOT / "Sinochem_Peer_Comparison.xlsx", data_only=True)
lay = json.loads((ROOT / "data" / "layout.json").read_text())
COLS = [("Revenue (USD bn)", "Rev $bn", "n1"), ("Revenue growth FY25", "Rev g", "p"), ("Gross margin", "GM", "p"),
        ("EBITDA margin", "EBITDA m", "p"), ("Net margin", "Net m", "p"), ("ROE (NI / avg equity)", "ROE", "p"),
        ("EBITDA / avg assets", "EBITDA/A", "p"), ("Net debt / EBITDA (x)", "ND/EBITDA", "x"), ("Interest cover (EBIT / interest, x)", "Int cover", "x"),
        ("Current ratio (x)", "Curr", "x"), ("OCF / EBITDA (cash conversion)", "OCF/EBITDA", "p"), ("FCF margin", "FCF m", "p")]


def fmt(v, k):
    if v is None: return ""
    if isinstance(v, str): return v
    return {"n1": f"{v:,.1f}", "p": f"{v*100:.1f}%", "x": f"{v:.1f}x"}[k]


def table(sector, stats=("Peer median - all peers",)):
    ws = wb[sector]
    hdr = {ws.cell(row=6, column=j).value: j for j in range(4, 40) if ws.cell(row=6, column=j).value}
    out = ["| Entity | " + " | ".join(c[1] for c in COLS) + " |", "|---|" + "---:|" * len(COLS)]
    for key, r in lay[sector]["rows"].items():
        name = ws.cell(row=r, column=1).value
        out.append(f"| {name} | " + " | ".join(fmt(ws.cell(row=r, column=hdr[c[0]]).value, c[2]) for c in COLS) + " |")
    for s in stats:
        r = lay[sector]["stats"][s]
        out.append(f"| *{s.replace('Peer median - ', 'Peer median: ')}* | " + " | ".join(fmt(ws.cell(row=r, column=hdr[c[0]]).value, c[2]) for c in COLS) + " |")
    return "\n".join(out)


def scorecard_summary():
    ws = wb["Scorecard"]
    rows, cur = {}, None
    for r in range(5, ws.max_row + 1):
        a, grp, met = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value
        if a:
            cur = a
            rows[cur] = dict(group=grp, tiers={"Top third": 0, "Mid-pack": 0, "Bottom third": 0}, m={})
        tier = ws.cell(row=r, column=9).value
        if tier:
            rows[cur]["tiers"]["Bottom third" if tier.startswith("Bottom") else tier] += 1
        rows[cur]["m"][met] = (ws.cell(row=r, column=4).value, ws.cell(row=r, column=5).value, ws.cell(row=r, column=7).value, ws.cell(row=r, column=8).value)
    out = ["| Sinochem entity | Compared with | Top third | Mid | Bottom third | EBITDA margin (vs median) | Net debt/EBITDA (vs median) | FCF margin (vs median) |", "|---|---|:---:|:---:|:---:|---|---|---|"]
    for name, d in rows.items():
        m = d["m"]
        e = m["EBITDA margin"]; n = m["Net debt / EBITDA (x)"]; f = m["FCF margin"]
        out.append(f"| {name} | {d['group']} | {d['tiers']['Top third']} | {d['tiers']['Mid-pack']} | {d['tiers']['Bottom third']} | "
                   f"{fmt(e[0], 'p')} ({fmt(e[1], 'p')}) | {fmt(n[0], 'x')} ({fmt(n[1], 'x')}) | {fmt(f[0], 'p')} ({fmt(f[1], 'p')}) |")
    return "\n".join(out)


def position_tables():
    ws = wb["Position"]
    blocks, cur = [], None
    for r in range(4, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        n = ws.cell(row=r, column=3).value
        if a and ws.cell(row=r, column=2).value is None and n is None and a.startswith(("ENERGY", "AGRICULTURE", "RUBBER")):
            cur = dict(title=a, rows=[]); blocks.append(cur)
        elif a and isinstance(n, int):
            g = lambda c: ws.cell(row=r, column=c).value  # noqa: E731
            rk = lambda c: f"{g(c)} of {n}" if isinstance(g(c), (int, float)) else "n/a"
            pct = g(15)
            cur["rows"].append(f"| {a} | {g(2)} (n={n}) | {g(4):,.1f} | {rk(5)} | {fmt(g(7), 'p')} ({rk(8)}) | {fmt(g(9), 'p')} ({rk(10)}) | {fmt(g(11), 'x')} ({rk(12)}) | "
                               f"{fmt(g(13), 'p')} ({rk(14)}) | **{g(16)}** ({pct * 100:.0f}%) | {g(17)} |")
    out = []
    for b in blocks:
        out.append(f"**{b['title']}**\n")
        out.append("| Sinochem entity | Compared with | Revenue (USD bn) | Size rank | EBITDA margin (rank) | ROE (rank) | Net debt/EBITDA (rank) | FCF margin (rank) | Operating position (percentile) | Market standing (public sources) |")
        out.append("|---|---|---:|:---:|---|---|---|---|---|---|")
        out.extend(b["rows"])
        out.append("")
    return "\n".join(out)


STATS = {"Energy": ("Peer median - all peers", "Peer median - State-owned majors", "Peer median - Private integrated refiners"),
         "Agriculture": ("Peer median - Crop-protection & seeds peers", "Peer median - Crop-nutrition peers"),
         "Rubber": ("Peer median - all peers", "Peer median - Global majors", "Peer median - China tyre producers")}

if __name__ == "__main__":
    for s, st in (("Energy", ("Peer median - all peers", "Peer median - State-owned majors", "Peer median - Private integrated refiners")),
                  ("Agriculture", ("Peer median - Crop-protection & seeds peers", "Peer median - Crop-nutrition peers")),
                  ("Rubber", ("Peer median - all peers", "Peer median - Global majors", "Peer median - China tyre producers"))):
        print(f"### {s}\n"); print(table(s, st)); print()
