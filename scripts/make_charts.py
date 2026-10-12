"""Peer-positioning charts, drawn from the values in the shipped workbook (emphasis form: Sinochem entity = accent, peers = grey)."""
import json
import pathlib
import re

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import openpyxl  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "charts"
OUT.mkdir(exist_ok=True)
wb = openpyxl.load_workbook(ROOT / "Sinochem_Peer_Comparison.xlsx", data_only=True)
lay = json.loads((ROOT / "data" / "layout.json").read_text())
D = json.loads((ROOT / "data" / "dataset.json").read_text())
ENT = {e["key"]: e for e in D["entities"]}

SURFACE, INK, INK2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df"
ACCENT, PEER = "#2a78d6", "#85847d"
PANELS = [("EBITDA margin", "EBITDA margin", "pct", "Higher is better"),
          ("EBITDA / avg assets", "Return on assets (EBITDA / avg total assets)", "pct", "Higher is better"),
          ("Net debt / EBITDA (x)", "Leverage (net debt / EBITDA, x)", "x", "Lower is better; negative = net cash")]
TITLES = {"Energy": "Energy: Sinochem Energy vs refining and petrochemical peers (FY2025)",
          "Agriculture": "Agriculture: Syngenta Group entities vs agri-input peers (FY2025)",
          "Rubber": "Rubber & tyres: Pirelli and Aeolus vs tyre and rubber peers (FY2025)"}
SHORT = lambda n: re.sub(r"\s*\(.*?\)|\s*[一-鿿]+", "", n).replace(" (global CP + Seeds)", "").strip()  # noqa: E731
SHORT_OVERRIDE = {"SYN": "Syngenta AG", "SLP": "Salt Lake Potash", "SSP": "Sinopec Shanghai", "SCE": "Sinochem Energy", "AEO": "Aeolus Tyre"}


def series(sector, header):
    ws = wb[sector]
    col = next(j for j in range(4, 40) if ws.cell(row=6, column=j).value == header)
    out = []
    for key, r in lay[sector]["rows"].items():
        v = ws.cell(row=r, column=col).value
        out.append((key, SHORT_OVERRIDE.get(key, SHORT(ENT[key]["name"])), ENT[key]["role"] == "Sinochem", v))
    return out


def draw(sector):
    n = len(lay[sector]["rows"])
    fig, axes = plt.subplots(1, 3, figsize=(15.5, 0.46 * n + 2.0), facecolor=SURFACE)
    fig.suptitle(TITLES[sector], x=0.012, ha="left", fontsize=14, fontweight="bold", color=INK)
    for ax, (hdr, ttl, kind, note) in zip(axes, PANELS):
        data = series(sector, hdr)
        num = [d for d in data if isinstance(d[3], (int, float))]
        asc = hdr.startswith("Net debt")
        num.sort(key=lambda d: d[3], reverse=not asc)
        nm = [d for d in data if not isinstance(d[3], (int, float))]
        rows = num + nm
        ypos = list(range(len(rows)))[::-1]
        ax.set_facecolor(SURFACE)
        vals = [d[3] for d in num]
        lo, hi = min(0, min(vals)), max(0, max(vals))
        pad = (hi - lo) * 0.22
        for y, (key, name, is_sin, v) in zip(ypos, rows):
            if isinstance(v, (int, float)):
                ax.barh(y, v, height=0.55, color=ACCENT if is_sin else PEER, zorder=3)
                txt = f"{v * 100:.1f}%" if kind == "pct" else f"{v:.1f}x"
                ax.text(v + (pad * 0.08 if v >= 0 else -pad * 0.08), y, txt, va="center", ha="left" if v >= 0 else "right", fontsize=9,
                        color=INK if is_sin else INK2, fontweight="bold" if is_sin else "normal", zorder=4)
            else:
                ax.text(pad * 0.08, y, "n.m. (negative EBITDA)" if v == "n.m." else "n/a", va="center", ha="left", fontsize=9, color=INK2, style="italic")
        ax.set_yticks(ypos)
        ax.set_yticklabels([r[1] for r in rows], fontsize=9.5)
        for lbl, (_, _, is_sin, _) in zip(ax.get_yticklabels(), rows):
            lbl.set_color(INK if is_sin else INK2)
            lbl.set_fontweight("bold" if is_sin else "normal")
        ax.set_xlim(lo - pad * (1.3 if lo < 0 else 0), hi + pad * 1.4)
        ax.xaxis.set_major_locator(matplotlib.ticker.MaxNLocator(nbins=6, steps=[1, 2, 5, 10]))
        ax.axvline(0, color=INK2, linewidth=0.8, zorder=2)
        ax.xaxis.grid(True, color=GRID, linewidth=0.7, zorder=0)
        ax.set_axisbelow(True)
        ax.tick_params(axis="x", labelsize=8.5, colors=INK2, length=0)
        ax.tick_params(axis="y", length=0)
        ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter((lambda v, _: f"{v * 100:.0f}%") if kind == "pct" else (lambda v, _: f"{v:.0f}x")))
        for s in ("top", "right", "left", "bottom"):
            ax.spines[s].set_visible(False)
        ax.set_title(ttl, loc="left", fontsize=10.5, color=INK, pad=17, fontweight="bold")
        ax.annotate(note, xy=(0, 1), xycoords="axes fraction", xytext=(0, 3), textcoords="offset points", fontsize=8.5, color=INK2, va="bottom")
    fig.text(0.012, 0.012, "Blue = Sinochem Holdings group entity; grey = listed peer. Reported-basis EBITDA (no add-back for impairments or one-offs). Source: company filings; see Sinochem_Peer_Comparison.xlsx.",
             fontsize=8.5, color=INK2)
    fig.tight_layout(rect=(0, 0.035, 1, 0.94), w_pad=2.2)
    path = OUT / f"{sector.lower()}_positioning.png"
    fig.savefig(path, dpi=170, facecolor=SURFACE)
    plt.close(fig)
    print("saved", path)


for s in ("Energy", "Agriculture", "Rubber"):
    draw(s)
